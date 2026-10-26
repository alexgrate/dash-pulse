from datetime import date, datetime
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 50_000


def wants_excel(request):
    return request.GET.get("format") == "xlsx"


def cell_value(value):
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    if isinstance(value, str) and len(value) in (10, 19) and value[4:5] == "-" and value[7:8] == "-":
        try:
            if len(value) == 10:
                return date.fromisoformat(value)
            if value[10] == "T":
                return datetime.fromisoformat(value)
        except ValueError:
            pass
    return value


def excel_bytes(columns, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for col, (label, _) in enumerate(columns, start=1):
        ws.cell(row=1, column=col, value=label).font = Font(bold=True)
    for r, row in enumerate(rows[:MAX_ROWS], start=2):
        for col, (_, key) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=col, value=cell_value(row.get(key)))
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for col, (label, key) in enumerate(columns, start=1):
        longest = max([len(label)] + [len(str(row.get(key) or "")) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(col)].width = min(longest + 2, 50)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_response(filename, columns, rows):
    response = HttpResponse(excel_bytes(columns, rows), content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response
