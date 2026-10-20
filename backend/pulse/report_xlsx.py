from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .report_data import change
from .report_pdf import DEFINITIONS

BRAND = "4A1A6B"
SOFT = "F4F0F8"
MUTED = "6B6880"
GOOD = "157A3C"
BAD = "B42318"

NAIRA = '"₦"#,##0.00'
NAIRA0 = '"₦"#,##0'
COUNT = "#,##0"
PCT = "0.0%"
DATE = "ddd dd mmm yyyy"

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
ZEBRA = PatternFill("solid", fgColor=SOFT)
LINE = Border(bottom=Side(style="thin", color="E2DDE9"))
FORMATS = {"count": COUNT, "money": NAIRA, "pct": PCT}


def sheet(wb, title, intro=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    if intro:
        ws["A2"] = intro
        ws["A2"].font = Font(italic=True, color=MUTED)
    return ws


def write_table(ws, top, headers, rows, formats, widths=None, freeze=True):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="left" if formats[c - 1] in (None, "text") else "right")
    for i, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=top + i, column=c, value=value)
            fmt = formats[c - 1]
            if fmt and fmt != "text":
                cell.number_format = fmt
            cell.border = LINE
            cell.alignment = Alignment(vertical="top", wrap_text=fmt == "text")
            if i % 2 == 0:
                cell.fill = ZEBRA
    ws.row_dimensions[top].height = 30
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"
    if freeze:
        ws.freeze_panes = ws.cell(row=top + 1, column=1)
    return top + len(rows)


def total_row(ws, row, values, formats):
    for c, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=value)
        cell.font = Font(bold=True)
        if formats[c - 1] and formats[c - 1] != "text":
            cell.number_format = formats[c - 1]
        cell.border = Border(top=Side(style="thin", color=BRAND))


def chart(ws, anchor, title, cats, series, stacked=False, y_title=None):
    c = BarChart()
    c.type = "col"
    c.title = title
    c.height, c.width = 8, 26
    c.legend.position = "b"
    if stacked:
        c.grouping = "stacked"
        c.overlap = 100
    for ref in series:
        c.add_data(ref, titles_from_data=True)
    c.set_categories(cats)
    if y_title:
        c.y_axis.title = y_title
    ws.add_chart(c, anchor)


def build_xlsx(report):
    cur, prev = report["current"], report["prior"]
    period, previous = report["period"], report["previous"]
    wb = Workbook()
    wb.remove(wb.active)

    ws = sheet(wb, "Summary")
    ws["A2"] = f"Dash MFB operations report · {period.label()} ({period.days} days)"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A3"] = (f"Compared with {previous.label()}. Generated {report['generated_at']:%d %b %Y %H:%M} WAT"
                f"{' by ' + report['generated_by'] if report['generated_by'] else ''}. Confidential, internal use only.")
    ws["A3"].font = Font(italic=True, color=MUTED)
    rows = []
    for r in report["comparison"]:
        rows.append([r["metric"], r["now"], r["before"], r["change"]])
    end = write_table(ws, 5, ["Measure", "This period", "Previous period", "Change"], rows,
                      ["text", None, None, None], [40, 22, 22, 14], freeze=False)
    for i, r in enumerate(report["comparison"], start=6):
        fmt = FORMATS[r["kind"]]
        ws.cell(row=i, column=2).number_format = fmt
        ws.cell(row=i, column=3).number_format = fmt
        cell = ws.cell(row=i, column=4)
        if not r.get("comparable", True):
            cell.value = "not comparable"
            cell.font = Font(italic=True, color=MUTED)
            cell.alignment = Alignment(horizontal="right")
            continue
        cell.number_format = '+0.0" pts";-0.0" pts"' if r["points"] else "+0.0%;-0.0%"
        if r["points"] and r["change"] is not None:
            cell.value = r["change"] * 100
        if r["change"] is not None and r["up_is_good"] is not None and abs(r["change"]) >= 0.005:
            good = (r["change"] > 0) == r["up_is_good"]
            cell.font = Font(color=GOOD if good else BAD, bold=True)
    row = end + 2
    ws.cell(row=row, column=1, value="Key findings").font = Font(bold=True, size=12, color=BRAND)
    for f in report["findings"]:
        row += 1
        ws.cell(row=row, column=1, value=f"• {f}")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
    if report["notes"]:
        row += 2
        ws.cell(row=row, column=1, value="Data notes").font = Font(bold=True, size=12, color=BRAND)
        for n in report["notes"]:
            row += 1
            ws.cell(row=row, column=1, value=f"• {n}")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 30

    ws = sheet(wb, "Daily", "One row per day (Lagos time). Totals at the bottom.")
    headers = ["Date", "Transactions", "Successful", "Failed", "Reversed", "Pending", "Success rate",
               "Money sent (₦)", "Money in (₦)", "Transfers received", "Sign-ups", "Accounts opened",
               "Received into other Dash accounts (₦)"]
    formats = [DATE, COUNT, COUNT, COUNT, COUNT, COUNT, PCT, NAIRA, NAIRA, COUNT, COUNT, COUNT, NAIRA]
    rows = [[s["date"], s["transactions"], s["success"], s["failed"], s["reversed"], s["pending"],
             s["success_rate"], s["value"], s["money_in"], s["money_in_count"], s["signups"], s["accounts"],
             s.get("money_in_other")]
            for s in cur["series"]]
    end = write_table(ws, 4, headers, rows, formats, [16, 13, 12, 10, 10, 10, 12, 18, 18, 12, 11, 12, 20])
    t = cur["totals"]
    total_row(ws, end + 1, ["Total", t["transactions"], t["success"], t["failed"], t["reversed"], t["pending"],
                            t["success_rate"], t["value"], t["money_in"], t["money_in_count"], t["signups"],
                            t["accounts"], t.get("money_in_other")], formats)
    cats = Reference(ws, min_col=1, min_row=5, max_row=end)
    chart(ws, f"O4", "Transactions per day", cats,
          [Reference(ws, min_col=3, min_row=4, max_row=end), Reference(ws, min_col=4, min_row=4, max_row=end)],
          stacked=True)
    chart(ws, f"O22", "Money sent and money in per day (₦)", cats,
          [Reference(ws, min_col=8, min_row=4, max_row=end), Reference(ws, min_col=9, min_row=4, max_row=end)])
    chart(ws, f"O40", "Sign-ups and accounts opened per day", cats,
          [Reference(ws, min_col=11, min_row=4, max_row=end), Reference(ws, min_col=12, min_row=4, max_row=end)])

    ws = sheet(wb, "Weekly", "Weeks run Monday to Sunday; the first and last weeks may be partial.")
    headers = ["Week starting", "Days in period", "Transactions", "Successful", "Failed", "Success rate",
               "Money sent (₦)", "Money in (₦)", "Sign-ups", "Accounts opened"]
    formats = [DATE, COUNT, COUNT, COUNT, COUNT, PCT, NAIRA, NAIRA, COUNT, COUNT]
    rows = [[w["date"], w["days"], w["transactions"], w["success"], w["failed"], w["success_rate"], w["value"],
             w["money_in"], w["signups"], w["accounts"]] for w in report["weekly"]]
    write_table(ws, 4, headers, rows, formats, [16, 10, 13, 12, 10, 12, 18, 18, 11, 12])

    ws = sheet(wb, "Products", "Every product customers used, busiest first, with the previous period for comparison.")
    before = {r["product"]: r for r in prev["products"]}
    headers = ["Product", "Transactions", "Share", "Successful", "Failed", "Reversed", "Failure rate",
               "Value sent (₦)", "Previous transactions", "Change", "Previous failure rate"]
    formats = ["text", COUNT, PCT, COUNT, COUNT, COUNT, PCT, NAIRA, COUNT, "+0.0%;-0.0%", PCT]
    rows = [[r["product"], r["count"], r["share"], r["success"], r["failed"], r["reversed"], r["failure_rate"],
             r["value"], before.get(r["product"], {}).get("count"),
             change(r["count"], before.get(r["product"], {}).get("count")),
             before.get(r["product"], {}).get("failure_rate")] for r in cur["products"]]
    write_table(ws, 4, headers, rows, formats, [28, 13, 9, 12, 10, 10, 12, 18, 14, 10, 14])

    ws = sheet(wb, "Failure reasons", "All failure reasons, most common first. 'Side' says who is usually responsible.")
    before = {r["reason"]: r["count"] for r in prev["reasons"]}
    headers = ["Reason", "Side", "Failures", "Share of failures", "Value (₦)", "Previous period", "Change",
               "Worst day", "Failures on worst day", "What it means"]
    formats = ["text", "text", COUNT, PCT, NAIRA, COUNT, "+0.0%;-0.0%", DATE, COUNT, "text"]
    rows = [[r["reason"], r["side"], r["count"], r["share"], r["value"], before.get(r["reason"]),
             change(r["count"], before.get(r["reason"])), r["worst_day"],
             r["days"].get(r["worst_day"]) if r["worst_day"] else None, r["hint"]] for r in cur["reasons"]]
    write_table(ws, 4, headers, rows, formats, [34, 18, 10, 12, 16, 12, 10, 16, 12, 60])

    ws = sheet(wb, "Receiving banks", "Transfers to other banks by receiving bank, busiest first.")
    headers = ["Receiving bank", "Transfers", "Successful", "Failed", "Failure rate", "Value sent (₦)"]
    formats = ["text", COUNT, COUNT, COUNT, PCT, NAIRA]
    rows = [[b["bank"], b["count"], b["success"], b["failed"], b["failure_rate"], b["value"]] for b in cur["banks"]]
    write_table(ws, 4, headers, rows, formats, [30, 12, 12, 10, 12, 18])

    ws = sheet(wb, "Onboarding", "People who started signing up in the period, and where they stopped.")
    end = write_table(ws, 4, ["Stage", "People", "% of sign-ups", "Dropped at this step"],
                      [[s["stage"], s["count"], s["share"], s["lost"]] for s in cur["funnel"]],
                      ["text", COUNT, PCT, COUNT], [34, 12, 14, 18], freeze=False)
    top = end + 3
    ws.cell(row=top - 1, column=1, value="Where people stopped").font = Font(bold=True, size=12, color=BRAND)
    write_table(ws, top, ["Last step reached", "People", "Share of sign-ups"],
                [[s["step"], s["count"], s["share"]] for s in cur["stuck"]], ["text", COUNT, PCT], freeze=False)

    ws = sheet(wb, "Rewards", "Rewards credited to customers in the period.")
    end = write_table(ws, 4, ["Reward", "Payouts", "Total paid (₦)"],
                      [[r["type"], r["count"], r["value"]] for r in cur["rewards"]],
                      ["text", COUNT, NAIRA], [30, 12, 18])
    total_row(ws, end + 1, ["Total", cur["totals"]["rewards_count"], cur["totals"]["rewards_value"]],
              ["text", COUNT, NAIRA])

    ws = sheet(wb, "Notable events", "Unusual days picked out automatically from the figures.")
    write_table(ws, 4, ["Date", "Event", "Detail"],
                [[d, kind, detail] for kind, d, detail in report["events"]], [DATE, "text", "text"], [18, 22, 80])

    ws = sheet(wb, "Definitions", "How each figure in this report is worked out.")
    write_table(ws, 4, ["Term", "Meaning"], [[t, m] for t, m in DEFINITIONS], ["text", "text"], [28, 100])

    for w in wb.worksheets:
        w.sheet_properties.tabColor = BRAND
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
