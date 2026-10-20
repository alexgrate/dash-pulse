import json
from copy import deepcopy
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import Client, SimpleTestCase, TestCase

from . import insights
from .reasons import classify

TYPICAL = {
    "weekday": "Friday",
    "days": 4,
    "hourly_median": [10] * 6 + [60] * 18,
    "hourly_low": [8] * 6 + [50] * 18,
    "hourly_high": [12] * 6 + [70] * 18,
    "so_far_median": 500,
    "total_median": 1000,
    "fraction_median": 0.5,
    "fraction_low": 0.45,
    "fraction_high": 0.55,
}

BASELINES = {
    "typical": TYPICAL,
    "reasons": {"Insufficient funds": 0.6, "Receiving bank is down": 0.1},
    "products": {"Utilities": {"count": 200, "failure_rate": 0.05, "reversal_rate": 0.08}},
    "liveness": 0.25,
    "rewards": {"per_tx": 1.0, "value_per_tx": 100},
    "signups": 50,
}

PAYLOAD = {
    "hour_now": 13,
    "transactions": {"today": {"count": 500, "failed": 50, "pending": 0}},
    "hourly": {"today": [10] * 6 + [60] * 18},
    "health": {
        "alert": None,
        "window_minutes": 15,
        "recent": {"count": 40, "failed": 6, "settled": 40, "failure_rate": 0.15},
        "baseline": {"failure_rate": 0.15},
        "reasons": [
            {"reason": "Insufficient funds", "kind": "customer", "today": 30, "recent": 2},
            {"reason": "Receiving bank is down", "kind": "system", "today": 5, "recent": 1},
        ],
        "products": [{"kind": "Utilities", "count": 40, "failure_rate": 0.05, "reversal_rate": 0.08}],
        "liveness": {"failure_rate": 0.25},
    },
    "funnel": {
        "days": 30,
        "stages": [
            {"label": "Signed up", "count": 100, "lost": 40},
            {"label": "Account opened", "count": 60, "lost": 0},
        ],
        "signups": {"today": 50, "yesterday": 50},
    },
    "geography": {"located": 0, "cities": [], "recent_minutes": 10},
    "rewards": {"today": {"count": 500, "value": 50000}, "last_paid_at": "2026-09-26T13:00:00"},
    "generated_at": "2026-09-26T14:00:00",
}


def run(**changes):
    p = deepcopy(PAYLOAD)
    for path, value in changes.items():
        if "__" not in path:
            p[path] = value
            continue
        target = p
        *keys, last = path.split("__")
        for k in keys:
            target = target[k]
        target[last] = value
    return insights.build(p, BASELINES)


def ids(result):
    return {i["id"] for i in result["items"]}


class InsightTests(SimpleTestCase):
    def test_normal_day_is_calm(self):
        result = run()
        severities = {i["severity"] for i in result["items"]}
        self.assertNotIn("bad", severities)
        self.assertNotIn("warn", severities)
        self.assertIn("volume", ids(result))

    def test_forecast_projects_from_typical_share_of_day(self):
        self.assertEqual(run()["forecast"]["projected"], 1000)

    def test_no_forecast_early_in_the_day(self):
        early = deepcopy(BASELINES)
        early["typical"] = {**TYPICAL, "fraction_median": 0.05}
        self.assertIsNone(insights.build(deepcopy(PAYLOAD), early)["forecast"])

    def test_busy_day_is_flagged(self):
        item = next(i for i in run(transactions__today={"count": 700, "failed": 50, "pending": 0})["items"]
                    if i["id"] == "volume")
        self.assertEqual(item["severity"], "good")
        self.assertIn("Busier", item["title"])

    def test_quiet_last_hour_warns(self):
        hourly = [10] * 6 + [60] * 18
        hourly[12] = 5
        item = next(i for i in run(hourly__today=hourly)["items"] if i["id"] == "hour")
        self.assertEqual(item["severity"], "warn")

    def test_alert_ranks_first(self):
        alert = {"title": "Transaction failures are spiking", "failure_rate": 0.4, "normal_rate": 0.15,
                 "failed": 20, "window_minutes": 15}
        first = run(health__alert=alert)["items"][0]
        self.assertEqual((first["id"], first["severity"]), ("alert", "bad"))

    def test_reason_shift_is_explained(self):
        reasons = [
            {"reason": "Insufficient funds", "kind": "customer", "today": 15, "recent": 1},
            {"reason": "Receiving bank is down", "kind": "system", "today": 25, "recent": 10},
        ]
        items = run(health__reasons=reasons)["items"]
        shift = next(i for i in items if i["id"] == "reason-Receiving bank is down")
        self.assertIn("NIP network", shift["detail"])

    def test_reversal_jump_is_flagged(self):
        products = [{"kind": "Utilities", "count": 40, "failure_rate": 0.05, "reversal_rate": 0.25}]
        self.assertIn("reversal-Utilities", ids(run(health__products=products)))

    def test_rewards_cost_growth_warns(self):
        item = next(i for i in run(rewards__today={"count": 700, "value": 70000})["items"] if i["id"] == "rewards")
        self.assertEqual(item["severity"], "warn")

    def test_without_history_it_still_answers(self):
        bare = {**BASELINES, "typical": None}
        result = insights.build(deepcopy(PAYLOAD), bare)
        self.assertIsNone(result["forecast"])
        self.assertIsNone(result["typical"])

    def test_account_numbers_never_reach_the_screen(self):
        for raw in ("Insufficient funds (0011058788)", "Account on PND, debit not allowed on this account: 0011230485",
                    "Something new happened for 0011230485"):
            label, _ = classify(raw)
            self.assertNotRegex(label, r"\d{6,}")

    def test_real_messages_are_grouped(self):
        self.assertEqual(classify("Insufficient Balance"), ("Insufficient funds", "customer"))
        self.assertEqual(classify("Your maximum debit limit per transaction is : 50,000.00")[1], "customer")
        self.assertEqual(classify("HTTP error occurred")[1], "system")
        self.assertEqual(classify(None)[1], "unknown")

    def test_pnd_failures_are_called_out(self):
        reasons = [
            {"reason": "Insufficient funds", "kind": "customer", "today": 20, "recent": 1},
            {"reason": "Customer's account is frozen", "kind": "account", "today": 15, "recent": 1},
        ]
        self.assertIn("reasons-pnd", ids(run(health__reasons=reasons)))

    def test_paused_rewards_are_explained(self):
        result = run(rewards__today={"count": 0, "value": 0}, rewards__last_paid_at="2026-09-20T18:00:00")
        item = next(i for i in result["items"] if i["id"] == "rewards-paused")
        self.assertIn("20 Sep", item["detail"])

    def test_signups_compare_with_typical_weekday(self):
        item = next(i for i in run(funnel__signups={"today": 20, "yesterday": 20})["items"] if i["id"] == "signups")
        self.assertIn("typical Friday", item["title"])

    def _banks(self, gt_recent_failed, gt_recent_system=None):
        def part(settled, failed, system):
            return {"count": settled, "settled": settled, "pending": 0, "failed": failed, "system_failed": system,
                    "failure_rate": failed / settled, "system_rate": system / settled}

        def bank(key, count, failed, system, r_settled, r_failed, r_system):
            return {"key": key, "bank": key.title(), "value": 1000.0, **part(count, failed, system),
                    "recent": part(r_settled, r_failed, r_system)}

        system = gt_recent_failed if gt_recent_system is None else gt_recent_system
        items = [bank("OPAY", 100, 15, 1, 20, 3, 0), bank("PALMPAY", 80, 12, 1, 15, 2, 0),
                 bank("GTBANK", 40, 8, 1, 12, gt_recent_failed, system)]
        for b in items:
            others = [x for x in items if x is not b]
            for field, key in (("failed", "others_rate"), ("system_failed", "others_system_rate")):
                b["recent"][key] = sum(x["recent"][field] for x in others) / sum(x["recent"]["settled"] for x in others)
            b["others_system_rate"] = sum(x["system_failed"] for x in others) / sum(x["settled"] for x in others)
        return {"items": items, "all": items, "bank_count": 3, "recent_minutes": 30,
                "total": {"count": 220, "value": 3000.0, "failure_rate": 35 / 220, "system_rate": 3 / 220}}

    def test_bank_outage_is_flagged(self):
        result = run(banks=self._banks(gt_recent_failed=10))
        item = next(i for i in result["items"] if i["id"] == "bank-GTBANK")
        self.assertEqual(item["severity"], "bad")
        self.assertIn("bank-side", item["detail"])

    def test_normal_banks_stay_quiet(self):
        self.assertNotIn("bank-GTBANK", ids(run(banks=self._banks(gt_recent_failed=2))))

    def test_customer_side_failures_never_flag_a_bank(self):
        self.assertNotIn("bank-GTBANK", ids(run(banks=self._banks(gt_recent_failed=6, gt_recent_system=0))))

    def test_bank_near_its_own_normal_is_not_flagged(self):
        base = deepcopy(BASELINES)
        base["banks"] = {"GTBANK": {"settled": 500, "failure_rate": 0.29, "system_rate": 0.3}}
        p = deepcopy(PAYLOAD)
        p["banks"] = self._banks(gt_recent_failed=5)
        self.assertNotIn("bank-GTBANK", {i["id"] for i in insights.build(p, base)["items"]})
        self.assertFalse(next(x for x in p["banks"]["all"] if x["key"] == "GTBANK")["failing_now"])

    def test_bank_far_above_its_own_normal_is_flagged(self):
        base = deepcopy(BASELINES)
        base["banks"] = {"GTBANK": {"settled": 500, "failure_rate": 0.2, "system_rate": 0.02}}
        p = deepcopy(PAYLOAD)
        p["banks"] = self._banks(gt_recent_failed=10)
        self.assertIn("bank-GTBANK", {i["id"] for i in insights.build(p, base)["items"]})

    def test_real_bank_names_are_cleaned(self):
        from .banks import canonical, key
        cases = {
            "OPAY": "OPay",
            "MONIEPOINT MICROFINANCE BANK": "Moniepoint MFB",
            "PALMPAY": "PalmPay",
            "GTBANK PLC": "GTBank",
            "SmartCash Payment Service bank": "SmartCash PSB",
            "MoMo PSB": "MoMo PSB",
            "FIRST BANK OF NIGERIA": "First Bank",
            "KUDA MICROFINANCE BANK": "Kuda MFB",
            "ACCESS BANK PLC": "Access Bank",
            None: "Unknown bank",
        }
        for raw, expected in cases.items():
            self.assertEqual(canonical(raw), expected)
        self.assertEqual(key("GTBANK PLC"), key("Guaranty Trust Bank"))


class LoginTests(TestCase):
    password = "correct-horse-battery-staple"

    def setUp(self):
        cache.clear()
        get_user_model().objects.create_user("ops", password=self.password)

    def post(self, client, url, body):
        return client.post(url, json.dumps(body), content_type="application/json")

    def test_dashboard_data_needs_login(self):
        self.assertEqual(self.client.get("/api/pulse").status_code, 401)

    def test_wrong_password_is_rejected_generically(self):
        r = self.post(self.client, "/api/auth/login", {"username": "ops", "password": "nope"})
        self.assertEqual(r.status_code, 401)
        r2 = self.post(self.client, "/api/auth/login", {"username": "nobody", "password": "nope"})
        self.assertEqual(r.json()["error"], r2.json()["error"])

    def test_lockout_after_repeated_failures(self):
        for _ in range(5):
            self.post(self.client, "/api/auth/login", {"username": "ops", "password": "nope"})
        r = self.post(self.client, "/api/auth/login", {"username": "ops", "password": self.password})
        self.assertEqual(r.status_code, 429)

    def test_login_then_data_then_logout(self):
        fake = {"clock_offset_seconds": 0}
        with patch("pulse.views.snapshot.build", return_value=fake):
            r = self.post(self.client, "/api/auth/login", {"username": "ops", "password": self.password})
            self.assertEqual(r.status_code, 200)
            self.assertTrue(r.json()["authenticated"])
            self.assertEqual(self.client.get("/api/pulse").status_code, 200)
            self.post(self.client, "/api/auth/logout", {})
            self.assertEqual(self.client.get("/api/pulse").status_code, 401)

    def test_remember_keeps_session_longer(self):
        self.post(self.client, "/api/auth/login", {"username": "ops", "password": self.password, "remember": True})
        self.assertGreater(self.client.session.get_expiry_age(), 7 * 24 * 3600)

    def test_login_requires_csrf_token(self):
        strict = Client(enforce_csrf_checks=True)
        r = self.post(strict, "/api/auth/login", {"username": "ops", "password": self.password})
        self.assertEqual(r.status_code, 403)
        strict.get("/api/auth/session")
        token = strict.cookies["csrftoken"].value
        r = strict.post("/api/auth/login", json.dumps({"username": "ops", "password": self.password}),
                        content_type="application/json", HTTP_X_CSRFTOKEN=token)
        self.assertEqual(r.status_code, 200)

    def test_manage_page_is_admin_only(self):
        r = self.client.get("/manage/auth/user/", follow=True)
        self.assertEqual(r.redirect_chain[-1][0], "/")
        self.post(self.client, "/api/auth/login", {"username": "ops", "password": self.password})
        self.assertNotEqual(self.client.get("/manage/auth/user/").status_code, 200)
        user = get_user_model().objects.get(username="ops")
        user.is_staff = user.is_superuser = True
        user.save()
        self.post(self.client, "/api/auth/login", {"username": "ops", "password": self.password})
        self.assertEqual(self.client.get("/manage/auth/user/").status_code, 200)
        self.assertTrue(self.client.get("/api/auth/session").json()["is_admin"])

    def test_admin_login_form_is_disabled(self):
        r = self.client.post("/manage/login/", {"username": "ops", "password": self.password})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(self.client.get("/api/auth/session").json()["authenticated"])


class BaselineMathTests(SimpleTestCase):
    def test_one_odd_day_does_not_move_typical(self):
        from .baselines import band, typical
        normal = [100, 104, 98, 102, 101, 99, 103]
        self.assertAlmostEqual(typical(normal + [900]), typical(normal), delta=2)
        low, high = band(normal + [900])
        self.assertLess(high, 110)
        self.assertGreater(low, 95)

    def test_few_days_still_work(self):
        from .baselines import band, typical
        self.assertEqual(typical([10, 20]), 15)
        self.assertEqual(band([10, 20]), (12.5, 17.5))


class TrendTests(SimpleTestCase):
    def summary(self, now, month_ago, last_week):
        return {"transactions": {"this_week": now, "last_week": last_week, "month_ago": month_ago,
                                 "vs_last_week": now / last_week - 1, "vs_month_ago": now / month_ago - 1}}

    def test_sustained_drop_is_flagged(self):
        result = run(trends={"summary": self.summary(667, 1000, 680)})
        item = next(i for i in result["items"] if i["id"] == "trend-transactions")
        self.assertEqual(item["severity"], "warn")
        self.assertIn("down 33%", item["title"])
        self.assertIn("two weeks", item["detail"])

    def test_small_change_is_ignored(self):
        self.assertNotIn("trend-transactions", ids(run(trends={"summary": self.summary(950, 1000, 960)})))

    def test_weekly_totals_from_series(self):
        from .trends import compare
        series = [{"transactions": 1000}] * 28 + [{"transactions": 700}] * 7
        s = compare(series, "transactions")
        self.assertEqual(s["this_week"], 700)
        self.assertAlmostEqual(s["vs_month_ago"], -0.3)


    def test_one_spike_day_does_not_move_the_week(self):
        from .trends import compare
        series = [{"transactions": 100}] * 34 + [{"transactions": 1000}]
        s = compare(series, "transactions")
        self.assertEqual(s["this_week"], 100)
        self.assertEqual(s["vs_month_ago"], 0)


class ExploreAccessTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import Group
        from django.http import JsonResponse

        from .auth import EXPLORER_GROUP, explorer_required

        self.group = Group.objects.get(name=EXPLORER_GROUP)
        self.user = get_user_model().objects.create_user("md", password="correct-horse-battery-staple")
        self.view = explorer_required(lambda request: JsonResponse({"ok": True}))

    def call(self):
        from django.test import RequestFactory

        request = RequestFactory().get("/api/explore/transactions?status=FAILED")
        request.user = self.user
        request.META["REMOTE_ADDR"] = "10.0.0.5"
        return self.view(request)

    def test_only_explorers_get_in_and_every_view_is_logged(self):
        from .models import ExploreLog

        self.assertEqual(self.call().status_code, 403)
        self.assertEqual(ExploreLog.objects.count(), 0)
        self.user.groups.add(self.group)
        self.assertEqual(self.call().status_code, 200)
        log = ExploreLog.objects.get()
        self.assertEqual((log.user, log.query, log.ip), (self.user, "status=FAILED", "10.0.0.5"))

    def test_people_list_needs_explore_and_merges_accounts(self):
        from datetime import datetime

        from django.test import RequestFactory

        from . import explore

        rows = [
            {"id": 1, "first": "Ada", "last": "Bello", "email": "Ada@x.com", "phone": "080", "phase": "ACCOUNT_CREATED",
             "started": datetime(2026, 10, 1, 9, 0), "opened": datetime(2026, 10, 1, 9, 6)},
            {"id": 2, "first": "Musa", "last": "Sani", "email": "musa@x.com", "phone": "081", "phase": "BVN_VALIDATION",
             "started": datetime(2026, 10, 1, 9, 0), "opened": None},
        ]
        found = [{"email": "ada@x.com", "account_number": "0011", "status": "ACTIVE"}]

        def call(query):
            request = RequestFactory().get(f"/api/explore/people?{query}")
            request.user = self.user
            with patch("pulse.explore.fetch", side_effect=[[{"n": 2}], [dict(r) for r in rows], found]):
                return explore.people(request)

        self.assertEqual(call("mode=signups&date=2026-10-01").status_code, 403)
        self.user.groups.add(self.group)
        body = json.loads(call("mode=signups&date=2026-10-01&stage=bogus").content)
        self.assertEqual((body["mode"], body["stage"], body["total"]), ("signups", None, 2))
        ada, musa = body["items"]
        self.assertEqual((ada["account"], ada["minutes_to_open"], ada["stage"]), ("0011", 6, "Account opened"))
        self.assertEqual((musa["account"], musa["done"], musa["stage"]), (None, False, "Stuck at BVN check"))
        self.assertEqual(json.loads(call("date=all").content)["date"], "all")

    def test_money_in_list_adds_account_names(self):
        from datetime import datetime

        from django.test import RequestFactory

        from . import explore

        self.user.groups.add(self.group)
        request = RequestFactory().get("/api/explore/money-in?date=2026-10-01")
        request.user = self.user
        rows = [{"provider_ref": "r1", "at": datetime(2026, 10, 1, 8, 0), "account": "0011", "amount": 5000,
                 "narration": "Inward transfer"}]
        side = [[{"n": 1, "value": 5000}], rows, [{"account_number": "0011", "account_name": "ADA BELLO"}]]
        with patch("pulse.explore.fetch", side_effect=side):
            body = json.loads(explore.money_in(request).content)
        self.assertEqual((body["total"], body["value"]), (1, 5000.0))
        self.assertEqual((body["items"][0]["name"], body["items"][0]["at"]), ("ADA BELLO", "2026-10-01T09:00:00"))

    def test_admin_alone_cannot_explore(self):
        self.user.is_staff = self.user.is_superuser = True
        self.user.save()
        self.assertEqual(self.call().status_code, 403)


class ExploreFilterTests(SimpleTestCase):
    row = {"id": 1, "at": "2026-10-04T10:15:00", "kind": "INTER", "amount": 5000.0, "account": "2000000001",
           "beneficiary": "Ada Bello", "bank": "OPay", "reference": "DMFB123", "outcome": "FAILED",
           "reason": "Insufficient funds", "reason_kind": "customer", "words": "dmfb123 2000000001 ada bello"}

    def q(self, **kw):
        base = {"status": None, "kinds": set(), "bank": None, "reason": None, "hour": None, "search": ""}
        return {**base, **kw}

    def test_filters_match_and_reject(self):
        from .explore import matches
        self.assertTrue(matches(self.row, self.q(status="FAILED", kinds={"INTER"}, bank="OPAY", hour=10)))
        self.assertTrue(matches(self.row, self.q(reason="Insufficient funds", search="ada")))
        self.assertFalse(matches(self.row, self.q(status="SUCCESS")))
        self.assertFalse(matches(self.row, self.q(kinds={"Airtime"})))
        self.assertFalse(matches(self.row, self.q(hour=11)))
        self.assertFalse(matches(self.row, self.q(bank="GTBank")))

    def test_search_ignores_spacing_case_and_word_order(self):
        from .explore import matches, words
        row = {**self.row, "words": words("DMFB9 0011273563 JAFAR  USMAN\u00a0 8020911442 Sebastine Agu")}
        for typed in ("JAFAR USMAN", "usman jafar", " jafar  ", "sebastine", "8020911442"):
            self.assertTrue(matches(row, self.q(search=words(typed))), typed)
        self.assertFalse(matches(row, self.q(search=words("jafar bello"))))


class PasswordEmailTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = get_user_model().objects.create_user("ops", email="ops@dash-mfb.com", password="old-Password-2026!")
        self.sent = []
        sender = patch("pulse.mailer.send", side_effect=lambda to, subject, html: self.sent.append((to, html)))
        sender.start()
        self.addCleanup(sender.stop)

    def post(self, url, body):
        return self.client.post(url, json.dumps(body), content_type="application/json")

    def link_parts(self):
        html = self.sent[-1][1]
        tail = html.split("#/set-password/")[1].split('"')[0]
        uid, token = tail.split("/")
        return {"uid": uid, "token": token}

    def test_forgot_sends_link_and_link_sets_password_once(self):
        r = self.post("/api/auth/forgot", {"email": "OPS@dash-mfb.com"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.sent[0][0], "ops@dash-mfb.com")
        parts = self.link_parts()
        self.assertEqual(self.post("/api/auth/check-link", parts).json(), {"valid": True, "username": "ops"})
        r = self.post("/api/auth/set-password", {**parts, "password": "short"})
        self.assertEqual(r.status_code, 400)
        r = self.post("/api/auth/set-password", {**parts, "password": "brand-New-Password-77"})
        self.assertEqual(r.status_code, 200)
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("brand-New-Password-77"))
        self.assertFalse(self.post("/api/auth/check-link", parts).json()["valid"])
        r = self.post("/api/auth/set-password", {**parts, "password": "another-New-Password-88"})
        self.assertEqual(r.status_code, 400)

    def test_forgot_does_not_reveal_unknown_emails(self):
        r = self.post("/api/auth/forgot", {"email": "nobody@dash-mfb.com"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json(), self.post("/api/auth/forgot", {"email": "ops@dash-mfb.com"}).json())
        self.assertEqual(len(self.sent), 1)

    def test_forgot_is_rate_limited(self):
        for _ in range(6):
            self.post("/api/auth/forgot", {"email": "ops@dash-mfb.com"})
        self.assertEqual(len(self.sent), 3)

    def test_disabled_users_get_nothing(self):
        self.user.is_active = False
        self.user.save()
        self.post("/api/auth/forgot", {"email": "ops@dash-mfb.com"})
        self.assertEqual(self.sent, [])

    def test_tampered_link_is_rejected(self):
        self.post("/api/auth/forgot", {"email": "ops@dash-mfb.com"})
        parts = self.link_parts()
        self.assertFalse(self.post("/api/auth/check-link", {**parts, "token": parts["token"] + "x"}).json()["valid"])
        self.assertFalse(self.post("/api/auth/check-link", {"uid": "!!", "token": "x"}).json()["valid"])

    def test_admin_creates_user_and_invite_is_sent(self):
        admin_user = get_user_model().objects.create_superuser("boss", password="boss-Password-2026!")
        self.client.force_login(admin_user)
        r = self.client.post(
            "/manage/auth/user/add/",
            {"username": "newbie", "email": "newbie@dash-mfb.com", "first_name": "New", "last_name": "", "explore": "on"},
        )
        self.assertEqual(r.status_code, 302)
        newbie = get_user_model().objects.get(username="newbie")
        self.assertFalse(newbie.has_usable_password())
        self.assertTrue(newbie.groups.filter(name="Explorer").exists())
        self.assertEqual(self.sent[-1][0], "newbie@dash-mfb.com")
        self.assertIn("newbie", self.sent[-1][1])

    def test_admin_toggles_explore_and_manage_from_edit_page(self):
        admin_user = get_user_model().objects.create_superuser("boss", password="boss-Password-2026!")
        self.client.force_login(admin_user)
        url = f"/manage/auth/user/{self.user.pk}/change/"
        self.assertContains(self.client.get(url), "Explore access")
        form = {"username": "ops", "email": "ops@dash-mfb.com", "first_name": "", "last_name": "", "is_active": "on"}
        self.client.post(url, {**form, "explore": "on", "is_staff": "on"})
        self.user.refresh_from_db()
        self.assertTrue(self.user.groups.filter(name="Explorer").exists())
        self.assertTrue(self.user.is_staff and self.user.is_superuser)
        self.client.post(url, form)
        self.user.refresh_from_db()
        self.assertFalse(self.user.groups.filter(name="Explorer").exists())
        self.assertFalse(self.user.is_staff or self.user.is_superuser)
        self.assertTrue(self.user.check_password("old-Password-2026!"))


class FailureMailTests(TestCase):
    def rows(self, *ids):
        from datetime import datetime

        return [{"id": i, "at": datetime(2026, 10, 12, 10, 0, i), "kind": "INTER", "amount": 1000 * i,
                 "account": "0011", "beneficiary": "Ada <b>", "bank": "OPAY", "reference": f"R{i}",
                 "session_id": f"S{i}", "provider_code": "91",
                 "message": "Insufficient funds (0011)" if i % 2 else "HTTP error occurred"} for i in ids]

    def run_with(self, rows):
        from .failure_mail import run_once

        sent = []
        with patch("pulse.failure_mail.recent_failures", return_value=rows), \
             patch("pulse.failure_mail.lagos_now"), \
             patch("pulse.failure_mail.mailer.send", side_effect=lambda to, s, h: sent.append((to, s, h))), \
             self.settings(PULSE_FAILURE_EMAILS=["ops@dash-mfb.com", "md@dash-mfb.com"]):
            run_once()
        return sent

    def test_first_run_skips_backlog_then_each_failure_is_mailed_once(self):
        self.assertEqual(self.run_with(self.rows(1, 2)), [])
        sent = self.run_with(self.rows(1, 2, 3, 4, 6))
        self.assertEqual(len(sent), 1)
        to, subject, html = sent[0]
        self.assertEqual(to, ["ops@dash-mfb.com", "md@dash-mfb.com"])
        self.assertIn("3 failed transactions", subject)
        self.assertIn("R3", html)
        self.assertIn("R4", html)
        self.assertIn("S6", html)
        self.assertIn("Insufficient funds", html)
        self.assertNotIn("R1", html)
        self.assertIn("Ada &lt;b&gt;", html)
        again = self.run_with(self.rows(1, 2, 3, 4, 6, 7))
        self.assertEqual(len(again), 1)
        self.assertIn("R7", again[0][2])
        self.assertNotIn("R6", again[0][2])
        self.assertEqual(self.run_with(self.rows(1, 2, 3, 4, 6, 7)), [])

    def test_failed_send_is_retried_next_time(self):
        from . import mailer
        from .failure_mail import run_once

        self.run_with(self.rows(2))
        with patch("pulse.failure_mail.recent_failures", return_value=self.rows(2, 4)), \
             patch("pulse.failure_mail.lagos_now"), \
             patch("pulse.failure_mail.mailer.send", side_effect=mailer.MailError("down")), \
             self.settings(PULSE_FAILURE_EMAILS=["ops@dash-mfb.com"]):
            with self.assertRaises(mailer.MailError):
                run_once()
        self.assertEqual(len(self.run_with(self.rows(2, 4))), 1)


class SignupHistoryTests(SimpleTestCase):
    def test_periods_start_on_monday_and_first_of_month(self):
        from datetime import date

        from .explore import history_periods, period_end, period_start

        wed = date(2026, 10, 14)
        self.assertEqual(period_start(wed, "week"), date(2026, 10, 12))
        self.assertEqual(period_end(date(2026, 10, 12), "week"), date(2026, 10, 18))
        self.assertEqual(period_end(date(2026, 2, 1), "month"), date(2026, 2, 28))
        self.assertEqual(period_end(date(2026, 12, 1), "month"), date(2026, 12, 31))
        months = history_periods(wed, "month")
        self.assertEqual((len(months), months[0], months[-1]), (12, date(2026, 10, 1), date(2025, 11, 1)))
        self.assertEqual(len(set(history_periods(wed, "day"))), 30)

    def test_range_end_is_clamped(self):
        from datetime import date

        from .explore import parse_range_end

        start = date(2026, 10, 1)
        self.assertEqual(parse_range_end(start, None), start)
        self.assertEqual(parse_range_end(start, "2026-09-01"), start)
        self.assertEqual(parse_range_end(start, "2026-10-07"), date(2026, 10, 7))
        self.assertEqual(parse_range_end(start, "2030-01-01"), date(2027, 11, 5))


def fake_report_period(start, end):
    from datetime import date, timedelta

    from .report_data import Period

    p = Period(date.fromisoformat(start), date.fromisoformat(end))
    series = []
    for i, d in enumerate(p.dates()):
        series.append({"date": d, "transactions": 100 + i, "success": 80, "failed": 15 + i % 3, "reversed": 2,
                       "pending": 3 + i, "settled": 97, "success_rate": 0.82, "value": 1_500_000.0 + i,
                       "money_in": 900_000.0, "money_in_count": 40, "signups": 60 + (200 if i == 3 else 0),
                       "finished": 20, "accounts": 0 if i == 5 else 22})
    totals = {"transactions": 1200, "success": 960, "failed": 190, "reversed": 24, "pending": 26, "settled": 1174,
              "success_rate": 0.8177, "failure_rate": 0.1618, "value": 18_000_000.0, "average_value": 18_750.0,
              "money_in": 10_800_000.0, "money_in_count": 480, "signups": 920, "accounts": 240,
              "finish_rate": 0.3, "active_users": 700, "rewards_count": 50, "rewards_value": 15_000.0}
    reasons = [{"reason": "Insufficient funds", "side": "Customer", "count": 120, "value": 900_000.0,
                "days": {p.start: 60, p.end: 60}, "share": 0.63, "hint": "Customers had no money.",
                "worst_day": p.start},
               {"reason": "Could not reach our <payment> partner", "side": "Bank or partner", "count": 70,
                "value": 400_000.0, "days": {p.start + timedelta(days=1): 70}, "share": 0.37, "hint": "",
                "worst_day": p.start + timedelta(days=1)}]
    return p, {
        "period": p, "series": series, "totals": totals,
        "products": [{"product": "Transfers to other banks", "count": 900, "success": 700, "failed": 170,
                      "reversed": 20, "value": 16_000_000.0, "failure_rate": 0.19, "share": 0.75}],
        "banks": [{"bank": "OPay", "count": 400, "success": 320, "failed": 70, "value": 6e6, "failure_rate": 0.18},
                  {"bank": "Bad Bank", "count": 50, "success": 30, "failed": 20, "value": 1e6, "failure_rate": 0.40}],
        "reasons": reasons,
        "funnel": [{"stage": "Signed up", "count": 920, "share": 1.0, "lost": 300},
                   {"stage": "Account opened", "count": 276, "share": 0.3, "lost": 0}],
        "stuck": [{"step": "Stuck at BVN check", "count": 100, "share": 0.11}],
        "rewards": [{"type": "Welcome bonus", "count": 50, "value": 15_000.0}],
    }


class ReportTests(TestCase):
    def setUp(self):
        from .report_data import Period

        self.admin = get_user_model().objects.create_superuser("boss", email="boss@dash-mfb.com",
                                                                password="boss-Password-2026!")
        fake = {}

        def collect(period):
            key = (period.start.isoformat(), period.end.isoformat())
            return fake.setdefault(key, fake_report_period(*key)[1])

        for target, value in (("pulse.report_data.collect", collect),
                              ("pulse.report_data.availability", lambda cur: ["A data note."]),
                              ("pulse.report_data.lagos_now", lambda: __import__("datetime").datetime(2026, 10, 20, 9))):
            patcher = patch(target, side_effect=value) if callable(value) else patch(target, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def post(self, **data):
        return self.client.post("/manage/reports/", {"period": "30", **data})

    def test_only_admins_can_open_reports(self):
        self.assertEqual(self.client.get("/manage/reports/", follow=True).redirect_chain[-1][0], "/")
        viewer = get_user_model().objects.create_user("viewer", password="viewer-Password-2026!")
        self.client.force_login(viewer)
        self.assertNotEqual(self.client.get("/manage/reports/").status_code, 200)
        self.client.force_login(self.admin)
        self.assertContains(self.client.get("/manage/reports/"), "Operations report")

    def test_presets_end_yesterday_and_custom_ranges_are_checked(self):
        from datetime import date

        from .report_data import resolve

        p = resolve("7", today=date(2026, 10, 20))
        self.assertEqual((p.start, p.end, p.days), (date(2026, 10, 13), date(2026, 10, 19), 7))
        self.assertEqual((p.previous().start, p.previous().end), (date(2026, 10, 6), date(2026, 10, 12)))
        p = resolve(None, "2026-10-10", "2026-10-01", today=date(2026, 10, 20))
        self.assertEqual((p.start, p.end), (date(2026, 10, 1), date(2026, 10, 10)))
        with self.assertRaises(ValueError):
            resolve(None, "2024-01-01", "2026-10-01", today=date(2026, 10, 20))

    def test_pdf_and_excel_downloads(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from .models import ReportLog

        self.client.force_login(self.admin)
        r = self.post(action="pdf")
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("2026-09-20_to_2026-10-19.pdf", r["Content-Disposition"])
        r = self.post(period="90", action="xlsx")
        wb = load_workbook(BytesIO(r.content))
        self.assertIn("Failure reasons", wb.sheetnames)
        self.assertEqual(wb["Daily"].cell(row=5, column=2).value, 100)
        self.assertEqual(list(ReportLog.objects.values_list("action", flat=True)), ["xlsx", "pdf"])

    def test_email_goes_only_to_company_addresses_with_both_files(self):
        from .models import ReportLog

        self.client.force_login(self.admin)
        sent = []
        with patch("pulse.report_views.start", side_effect=lambda fn, *a: fn(*a)), \
             patch("pulse.report_views.mailer.send", side_effect=lambda to, s, h, att: sent.append((to, s, att))):
            r = self.post(action="email", email="someone@gmail.com")
            self.assertContains(r, "can only be sent to @dash-mfb.com")
            self.assertEqual(sent, [])
            self.post(action="email", email="MD@Dash-MFB.com")
        to, subject, attachments = sent[0]
        self.assertEqual(to, ["md@dash-mfb.com"])
        self.assertIn("20 Sep 2026", subject)
        self.assertEqual([a[0].rsplit(".", 1)[1] for a in attachments], ["pdf", "xlsx"])
        self.assertTrue(attachments[0][2].startswith(b"%PDF"))
        self.assertEqual(ReportLog.objects.get().status, "sent")

    def test_findings_and_events_read_the_figures(self):
        from .report_data import build, resolve

        report = build(resolve("30"), "boss")
        text = " ".join(report["findings"])
        self.assertIn("Insufficient funds", text)
        kinds = {e[0] for e in report["events"]}
        self.assertIn("No accounts opened", kinds)
        self.assertIn("Sign-up spike", kinds)
        self.assertEqual(report["notes"], ["A data note."])
